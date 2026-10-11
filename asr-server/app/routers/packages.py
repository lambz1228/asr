"""任务（文本包）导入与下发。

支持的提交方式：
1) multipart/form-data：上传 .xlsx / .csv / .json 文件（字段名 file），可附带表单参数。
2) application/json：直接提交 JSON 对象。

文件格式（推荐：导出的 Excel 两列表）
    序号 | 汉语文本
    1    | 今天有空吗？我们一起去街上逛逛，买点好吃的回来。
    2    | 手机没电了，我先充一会儿电，等下再和你聊天。
  - 「序号」列自动忽略（不进 textId，避免跨任务包冲突）
  - 「汉语文本」列作为待录文本；同义表头：文本 / 内容 / 中文文本 / 文本内容 / 待录制文本 / 语句
  - 无表头时按「序号 + 文本」两列兜底识别

下发方式（distribute）：
  - 不传 assignees            → 进入公共池，采集员在「我的任务」页自行领取
  - assignees + distribute=all   → 每个采集员各拿到全部文本（同一文本多人录制）
  - assignees + distribute=split → 文本按顺序轮流分摊给各采集员
"""
import csv
import io
import json
import uuid

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from sqlalchemy import func, select

from .. import models
from ..deps import require_admin, get_db, record_audit
from ..responses import ok

router = APIRouter()

# 表头别名 → 标准字段（统一小写、去空格后匹配）。
# 注意：「序号」等列不在表中，解析时会被自动忽略。
_ALIASES = {
    "text": "text", "文本": "text", "内容": "text", "待录制文本": "text",
    "汉语文本": "text", "中文文本": "text", "文本内容": "text",
    "语句": "text", "句子": "text", "sentence": "text", "content": "text",
    "textid": "textId", "text_id": "textId", "文本id": "textId",
    "assignee": "assignee", "采集员": "assignee", "分配": "assignee",
    "采集员id": "assignee", "用户": "assignee", "账号": "assignee",
}

TEMPLATE_CSV = (
    "序号,汉语文本\n"
    "1,今天有空吗？我们一起去街上逛逛，买点好吃的回来。\n"
    "2,手机没电了，我先充一会儿电，等下再和你聊天。\n"
    "3,晚上想看个电视吗？我们找好看的节目慢慢看。\n"
)


def _norm_key(k: str) -> str | None:
    return _ALIASES.get((k or "").strip().lower().replace(" ", ""))


def _is_seq(s: str) -> bool:
    """判断是否为「序号」列（纯数字且不过长）。"""
    s = (s or "").strip()
    return s.isdigit() and len(s) <= 6


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        out: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            out.append(["" if v is None else str(v).strip() for v in row])
        return out
    finally:
        wb.close()


def _json_to_rows(raw: str) -> list[dict]:
    data = json.loads(raw)
    if isinstance(data, dict):
        data = data.get("texts", [])
    rows = []
    for it in data:
        if isinstance(it, str):
            rows.append({"text": it})
        else:
            rows.append({
                "text": it.get("text") or it.get("content") or "",
                "textId": it.get("textId") or it.get("text_id") or "",
                "assignee": it.get("assignee") or "",
            })
    return rows


def _table_to_rows(table: list[list[str]]) -> list[dict]:
    """把二维表（xlsx/csv）转为 [{text, textId?, assignee?}]。"""
    table = [list(r) for r in table if any((c or "").strip() for c in r)]
    if not table:
        return []
    header = [(c or "").strip() for c in table[0]]
    mapped = {i: _norm_key(c) for i, c in enumerate(header)}
    has_header = "text" in mapped.values()
    body = table[1:] if has_header else table

    rows: list[dict] = []
    for r in body:
        cells = [(c or "").strip() for c in r]
        row: dict[str, str] = {}
        if has_header:
            for i, key in mapped.items():
                if key and i < len(cells):
                    row[key] = cells[i]
        else:
            # 无表头兜底：识别「序号 + 文本」结构；序号只用于识别，不作为 textId
            if len(cells) >= 2 and _is_seq(cells[0]):
                row["text"] = cells[1]
            else:
                row["text"] = next((c for c in cells if c), "")
        rows.append(row)
    return rows


def _load_rows(data: bytes, filename: str = "") -> list[dict]:
    """按文件类型（xlsx / csv / json）统一解析为待录文本行。"""
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ValueError("不支持旧版 .xls，请另存为 .xlsx 或 .csv 后重试")
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return _table_to_rows(_read_xlsx(data))
    raw = data.decode("utf-8-sig", errors="replace").strip()
    if not raw:
        return []
    if raw[0] in "[{":
        return _json_to_rows(raw)
    return _table_to_rows([r for r in csv.reader(io.StringIO(raw))])


def _split_list(v) -> list[str]:
    """兼容 'a,b' / ['a','b'] / None 三种写法，返回去重后的有序列表。"""
    if v is None:
        return []
    if isinstance(v, (list, tuple)):
        parts = [str(x).strip() for x in v]
    else:
        parts = [p.strip() for p in str(v).replace("，", ",").split(",")]
    seen, out = set(), []
    for p in parts:
        if p and p not in seen:
            seen.add(p)
            out.append(p)
    return out


@router.get("/admin/packages")
async def list_packages(principal=Depends(require_admin), db=Depends(get_db)):
    """列出任务包及其文本/任务数量。"""
    pkgs = (
        await db.execute(select(models.TaskPackage).order_by(models.TaskPackage.id))
    ).scalars().all()
    items = []
    for p in pkgs:
        n_texts = await db.scalar(
            select(func.count()).select_from(models.TextItem).where(
                models.TextItem.pkg_id == p.pkg_id
            )
        )
        n_asg = await db.scalar(
            select(func.count()).select_from(models.TaskAssignment).where(
                models.TaskAssignment.pkg_id == p.pkg_id
            )
        )
        n_pending = await db.scalar(
            select(func.count()).select_from(models.TaskAssignment).where(
                models.TaskAssignment.pkg_id == p.pkg_id,
                models.TaskAssignment.status == "pending",
            )
        )
        n_pool = await db.scalar(
            select(func.count()).select_from(models.TaskAssignment).where(
                models.TaskAssignment.pkg_id == p.pkg_id,
                models.TaskAssignment.assignee.is_(None),
            )
        )
        assignees = (
            await db.execute(
                select(models.TaskAssignment.assignee)
                .where(
                    models.TaskAssignment.pkg_id == p.pkg_id,
                    models.TaskAssignment.assignee.is_not(None),
                )
                .distinct()
            )
        ).scalars().all()
        items.append({
            "pkgId": p.pkg_id,
            "name": p.name,
            "params": p.params,
            "paramVersion": p.param_version,
            "createdBy": p.created_by,
            "texts": n_texts,
            "assignments": n_asg,
            "pending": n_pending,
            "pool": n_pool,
            "assignees": list(assignees),
        })
    return ok({"list": items})


@router.get("/admin/packages/template")
async def download_template(principal=Depends(require_admin)):
    """下载 CSV 导入模板（UTF-8 BOM，Excel 可直接打开）。"""
    return Response(
        content="\ufeff" + TEMPLATE_CSV,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=task_import_template.csv"},
    )


@router.post("/admin/packages/import")
async def import_package(
    request: Request, principal=Depends(require_admin), db=Depends(get_db)
):
    """导入任务包并按需下发给采集员。

    multipart/form-data 字段：
      file        .xlsx / .csv / .json 文件（必填）
      pkgId       任务包ID（可选，缺省自动生成 pkg_xxxxxxxx）
      name        任务包名称（可选，缺省用 pkgId）
      assignees   下发对象：采集员 user_id，多个用逗号分隔（可选，留空=公共池）
      distribute  下发方式：all=每人一份全部文本（默认）| split=轮流分摊
      assignee    兼容旧版的单个采集员（可选）
      sampleRate  采样率（可选，默认 16000）
      channels    声道数（可选，默认 1）
      format      采集格式（可选，默认 aac）

    application/json 字段：
      {pkgId, name, assignees:[...], distribute:"all",
       params:{sample_rate,channels,format}, texts:[{text,textId,assignee}]}
    """
    ctype = (request.headers.get("content-type") or "").lower()
    pkg_id = name = ""
    assignees: list[str] = []
    distribute = "all"
    sample_rate, channels, audio_format = 16000, 1, "aac"
    rows: list[dict] = []

    if "multipart/form-data" in ctype:
        form = await request.form()
        f = form.get("file")
        pkg_id = (form.get("pkgId") or "").strip()
        name = (form.get("name") or "").strip()
        assignees = _split_list(form.get("assignees")) or _split_list(form.get("assignee"))
        distribute = (form.get("distribute") or "all").strip().lower()
        try:
            sample_rate = int(form.get("sampleRate") or 16000)
            channels = int(form.get("channels") or 1)
        except (TypeError, ValueError):
            return {"code": 1, "message": "sampleRate/channels 需为整数", "data": None}
        audio_format = (form.get("format") or "aac").strip()
        if f is None or not hasattr(f, "read"):
            return {"code": 1, "message": "缺少文件字段 file", "data": None}
        try:
            rows = _load_rows(await f.read(), getattr(f, "filename", "") or "")
        except ValueError as e:
            return {"code": 1, "message": str(e), "data": None}
        except json.JSONDecodeError as e:
            return {"code": 1, "message": f"文件解析失败：{e}", "data": None}
        except Exception as e:  # 损坏的 xlsx 等
            return {"code": 1, "message": f"文件解析失败：{e}", "data": None}
    else:
        try:
            body = await request.json()
        except Exception:
            return {"code": 1, "message": "请求体需为 JSON 或 multipart/form-data", "data": None}
        pkg_id = (body.get("pkgId") or "").strip()
        name = (body.get("name") or "").strip()
        assignees = _split_list(body.get("assignees")) or _split_list(body.get("assignee"))
        distribute = (body.get("distribute") or "all").strip().lower()
        params = body.get("params") or {}
        sample_rate = int(params.get("sample_rate", 16000))
        channels = int(params.get("channels", 1))
        audio_format = params.get("format", "aac")
        rows = _json_to_rows(json.dumps(body.get("texts", []), ensure_ascii=False))

    if distribute not in ("all", "split"):
        return {"code": 1, "message": "distribute 仅支持 all 或 split", "data": None}
    if not rows:
        return {"code": 1, "message": "未解析到任何文本行（请检查表头/内容）", "data": None}

    pkg_id = pkg_id or ("pkg_" + uuid.uuid4().hex[:8])
    if await db.scalar(
        select(models.TaskPackage).where(models.TaskPackage.pkg_id == pkg_id)
    ):
        return {"code": 1, "message": f"任务包已存在：{pkg_id}", "data": None}

    name = name or pkg_id
    db.add(models.TaskPackage(
        pkg_id=pkg_id,
        name=name,
        params={
            "sample_rate": sample_rate,
            "channels": channels,
            "format": audio_format,
            "assignees": assignees,
            "distribute": distribute,
        },
        created_by=principal["uid"],
    ))
    await db.flush()

    created_texts = created_asg = skipped = 0
    errors: list[str] = []
    idx = 0
    for i, r in enumerate(rows, 1):
        text = (r.get("text") or "").strip()
        if not text:
            skipped += 1
            errors.append(f"第 {i} 行：文本为空，已跳过")
            continue
        tid = (r.get("textId") or "").strip() or f"{pkg_id}_t{i:04d}"
        if await db.scalar(select(models.TextItem).where(models.TextItem.text_id == tid)):
            skipped += 1
            errors.append(f"第 {i} 行：textId 已存在（{tid}），已跳过")
            continue
        db.add(models.TextItem(text_id=tid, pkg_id=pkg_id, text=text, text_version=1))
        created_texts += 1

        # 决定该文本下发给谁：行内指定 > 统一下发 > 公共池
        row_assignee = (r.get("assignee") or "").strip()
        if row_assignee:
            who_list: list[str | None] = [row_assignee]
        elif assignees:
            who_list = list(assignees) if distribute == "all" else [assignees[(i - 1) % len(assignees)]]
        else:
            who_list = [None]
        for who in who_list:
            idx += 1
            db.add(models.TaskAssignment(
                assignment_id=f"{pkg_id}_a{idx:04d}",
                pkg_id=pkg_id,
                text_id=tid,
                assignee=who,
                status="pending",
            ))
            created_asg += 1

    record_audit(
        db, principal["uid"], principal["role"], "import_package", "package", pkg_id,
        {
            "texts": created_texts,
            "assignments": created_asg,
            "skipped": skipped,
            "assignees": assignees,
            "distribute": distribute,
        },
    )
    await db.commit()
    return ok({
        "pkgId": pkg_id,
        "name": name,
        "params": {"sampleRate": sample_rate, "channels": channels, "format": audio_format},
        "assignees": assignees,
        "distribute": distribute,
        "createdTexts": created_texts,
        "createdAssignments": created_asg,
        "skipped": skipped,
        "errors": errors[:50],
    })
